#!/usr/bin/env python3
"""生成 PPT 智能体的测试目录。

用法: python3 scripts/make_test_ppts.py test-ppts [--big-mb 25]

生成的文件覆盖以下场景：
  * 中英文混排、表格、图表、组合形状、隐藏页、无标题页、多段备注
  * 内容相近的 白底 / 黑底 两个版本（用于底色切换提示与相似分组）
  * 版本副本（v2）与完全相同的重复文件（用于"已解析过"去重）
  * 嵌入视频 + 字幕（mp4 内封 mov_text 字幕、pptx 内 vtt 字幕部件）的大文件
  * 外链视频 + 同名 .srt 字幕
  * 母版级深色背景、整页图片背景
  * 旧版 .ppt（需要 LibreOffice 转换）、空白演示、损坏文件、Office 锁文件 ~$
依赖: python-pptx, Pillow；可选 ffmpeg（生成视频）、soffice（生成 .ppt）
"""
import argparse
import io
import os
import random
import shutil
import subprocess
import sys
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
DARK = RGBColor(0x11, 0x14, 0x1A)
ACCENT = RGBColor(0xC0, 0x1E, 0x2E)

TURNAROUND = [
    ("航班过站保障 Aircraft Turnaround 概述",
     ["过站保障（Turnaround）是飞机从落地上轮挡到下一次撤轮挡之间的全部地面作业",
      "涉及加油、配餐、清洁、行李装卸、旅客上下机等 20 余个保障节点",
      "Target: 窄体机过站时间压缩到 35 分钟以内"],
     "各位好，今天我们介绍航班过站保障。过站时间直接影响航班准点率和飞机利用率，"
     "是机场和航司共同关注的核心指标。\n请注意窄体机 35 分钟的目标。"),
    ("关键里程碑 Key Milestones",
     ["AIBT 实际上轮挡时间 (Actual In-Block Time)",
      "TOBT 目标撤轮挡时间 (Target Off-Block Time)",
      "TSAT 目标启动许可时间 (Target Start-up Approval Time)",
      "AOBT 实际撤轮挡时间"],
     "这一页讲 A-CDM 里与过站相关的四个里程碑，TOBT 由航司或地服负责更新，TSAT 由塔台排序给出。"),
    ("保障流程甘特图 Turnaround Gantt", None,
     "甘特图展示了关键路径：下客 → 清洁 → 上客，加油可以与配餐并行，但有旅客在机上时加油需要消防监护。"),
    ("地面保障设备 GSE",
     ["客梯车 / 廊桥 Passenger Boarding Bridge",
      "行李传送带车 Belt Loader",
      "牵引车 Pushback Tug",
      "电源车 GPU 与空调车 PCA"],
     "GSE 的调度效率决定了保障节点能否按时开始，我们通过物联网定位实现车辆实时调度。"),
    ("过站时间统计 Turnaround KPI", "TABLE",
     "表格对比了三个基地的平均过站时间，北京基地通过协同决策把平均时间缩短了 6 分钟。"),
    ("准点率趋势 OTP Trend", "CHART",
     "图表显示实施 A-CDM 后放行正常率逐月提升，九月达到 91%。"),
    ("优化建议 Recommendations",
     ["推行 TOBT 准确率考核，误差控制在 ±5 分钟",
      "关键路径节点数字化签到，减少人工通报",
      "引入 AI 预测模型预估过站时间 Predictive Turnaround"],
     "最后是三条优化建议，其中 AI 预测模型已经在试点，预计明年推广。谢谢大家。"),
]

ACDM = [
    ("A-CDM 机场协同决策介绍",
     ["Airport Collaborative Decision Making", "机场、航司、空管、地服信息共享", "提升航班放行正常率与资源利用率"],
     "A-CDM 是欧洲航管组织提出的理念，国内已在主要枢纽机场推广。"),
    ("信息共享 Information Sharing",
     ["统一的航班运行数据平台", "里程碑事件实时推送", "各方共用同一时间基准"],
     "信息共享是 A-CDM 的基础，没有统一数据平台就谈不上协同。"),
    ("可变滑行时间 Variable Taxi Time",
     ["按机位与跑道组合计算 EXOT / EXIT", "结合历史数据与实时场面监视"],
     "滑行时间预测越准，TSAT 排序越合理。"),
    ("除冰协同 De-icing", ["冬季除冰排队纳入 CDM 排序", "除冰坪资源分配"],
     "冬季运行时除冰是影响放行的重要因素。"),
]

AI_PLATFORM = [
    ("Smart Airport AI Platform",
     ["Computer vision for apron monitoring", "Predictive turnaround with machine learning", "Large language model assistant for operations"],
     "Welcome. This deck introduces our smart airport AI platform and its three pillars."),
    ("Apron Vision 机坪视觉识别",
     ["Detect GSE arrival and departure events automatically", "Auto-capture turnaround milestones", "Safety: FOD and zone intrusion alerts"],
     "Cameras on the apron recognise when the fuel truck or catering truck arrives, so milestones are recorded without manual input."),
    ("Architecture 平台架构", "GROUP",
     "Edge nodes run the vision models, the cloud hosts training and the RAG knowledge base."),
    ("RAG Knowledge Assistant",
     ["Hybrid retrieval: BM25 + dense embeddings", "Reranking with cross-encoder", "Grounded answers with citations"],
     "The assistant answers SOP questions using retrieval augmented generation, always citing the source page."),
    ("Roadmap 2027", ["Q1 pilot at Terminal 3", "Q2 rollout to all stands", "Q3 digital twin integration"],
     "Our roadmap for next year focuses on scaling the pilot."),
]

CLOUD = [
    ("Cloud Architecture Overview", ["Multi-region active-active deployment", "Kubernetes based microservices", "Zero trust security"],
     "This overview covers our target cloud architecture."),
    ("Data Platform", ["Lakehouse with streaming ingestion", "Unified metadata catalog", "Vector database for semantic search"],
     "The data platform supports both BI and AI workloads."),
    ("Cost Optimization", ["Reserved instances and spot fleets", "FinOps dashboards", "Auto scaling policies"],
     "Cost optimization saved 23 percent last year."),
]

FINANCE = [
    ("2026 年第三季度经营分析", ["营业收入同比增长 12%", "毛利率 31.5%", "经营现金流持续为正"],
     "本季度整体经营情况良好，下面分业务线展开。"),
    ("分业务收入", "TABLE2", "地面服务收入占比最高，货运业务增长最快。"),
    ("成本控制措施", ["集中采购降低燃油成本", "人力排班优化", "设备共享减少闲置"],
     "成本控制方面我们重点抓了三件事。"),
]


def set_bg(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def style_text(shape, color, size=None):
    if not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = color
            if size:
                r.font.size = Pt(size)


def add_table(slide, rows, color=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.8), Inches(1.8), Inches(8.4), Inches(0.5) * len(rows))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = shape.table.cell(i, j)
            cell.text = str(val)
            if color is not None:
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = color
    return shape


def add_chart(slide):
    data = CategoryChartData()
    data.categories = ["5月", "6月", "7月", "8月", "9月"]
    data.add_series("放行正常率 OTP %", (82, 84, 87, 89, 91))
    slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(1), Inches(1.6), Inches(8), Inches(4.5), data)


def add_gantt(slide, color):
    tasks = [("下客 Deplaning", 0, 8), ("清洁 Cleaning", 8, 10), ("配餐 Catering", 5, 15), ("加油 Fueling", 6, 14),
             ("行李装卸 Baggage", 2, 25), ("上客 Boarding", 18, 15)]
    for i, (name, start, dur) in enumerate(tasks):
        top = Inches(1.6 + i * 0.7)
        lbl = slide.shapes.add_textbox(Inches(0.5), top, Inches(2.4), Inches(0.5))
        lbl.text_frame.text = name
        style_text(lbl, color, 14)
        bar = slide.shapes.add_shape(1, Inches(3 + start * 0.17), top, Inches(dur * 0.17), Inches(0.45))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.text_frame.text = f"{dur} min"


def add_group(slide, color):
    grp = slide.shapes.add_group_shape()
    for i, name in enumerate(["Edge 边缘节点", "Cloud 云端训练", "RAG 知识库"]):
        box = grp.shapes.add_shape(1, Inches(0.8 + i * 3), Inches(2.5), Inches(2.6), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0x2A, 0x4D, 0x8F)
        box.text_frame.text = name
        style_text(box, WHITE, 18)


def build_deck(path, slides_spec, dark=False, master_dark=False, title_prefix="", hidden=(), extra=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    prs.core_properties.author = "测试作者"
    prs.core_properties.title = os.path.splitext(os.path.basename(path))[0]
    if master_dark:
        # 母版背景设为深色，幻灯片本身不设背景（测试继承）
        master = prs.slide_master
        csld = master.element.find(qn("p:cSld"))
        for old in csld.findall(qn("p:bg")):
            csld.remove(old)
        bg = etree.Element(qn("p:bg"))
        csld.insert(0, bg)
        bgpr = etree.SubElement(bg, qn("p:bgPr"))
        sf = etree.SubElement(bgpr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val="0B0F1A")
        etree.SubElement(bgpr, qn("a:effectLst"))
    fg = WHITE if (dark or master_dark) else BLACK
    for idx, (title, body, notes) in enumerate(slides_spec):
        layout = prs.slide_layouts[0] if idx == 0 else prs.slide_layouts[1 if isinstance(body, list) else 5]
        s = prs.slides.add_slide(layout)
        if dark:
            set_bg(s, DARK)
        s.shapes.title.text = title_prefix + title
        style_text(s.shapes.title, fg)
        if isinstance(body, list):
            ph = s.placeholders[1]
            tf = ph.text_frame
            tf.text = body[0]
            for line in body[1:]:
                p = tf.add_paragraph()
                p.text = line
                if line.startswith("TOBT") or line.startswith("TSAT"):
                    p.level = 1
            style_text(ph, fg)
        elif body == "TABLE":
            add_table(s, [["基地 Base", "平均过站(分钟)", "TOBT 准确率"], ["北京", "38", "86%"], ["上海", "41", "81%"], ["广州", "44", "78%"]],
                      fg if dark else None)
        elif body == "TABLE2":
            add_table(s, [["业务线", "收入(亿元)", "同比"], ["地面服务", "12.4", "+8%"], ["货运", "6.1", "+21%"], ["商业零售", "3.3", "+5%"]])
        elif body == "CHART":
            add_chart(s)
        elif body is None:
            add_gantt(s, fg)
        elif body == "GROUP":
            add_group(s, fg)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        if idx in hidden:
            s._element.set("show", "0")
    if extra:
        extra(prs, fg)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    prs.save(path)
    return path


def make_video(path, seconds, size, mbps, subtitles):
    """生成带 mov_text 内封字幕的 mp4。"""
    srt = path + ".tmp.srt"
    with open(srt, "w", encoding="utf-8") as f:
        f.write(subtitles)
    w, h = size
    cmd = ["ffmpeg", "-y", "-loglevel", "error",
           "-f", "lavfi", "-i", f"testsrc2=size={w}x{h}:rate=25",
           "-f", "lavfi", "-i", "sine=frequency=440",
           "-i", srt, "-t", str(seconds),
           "-map", "0:v", "-map", "1:a", "-map", "2:s",
           "-c:v", "libx264", "-preset", "ultrafast", "-b:v", f"{mbps}M", "-minrate", f"{mbps}M", "-maxrate", f"{mbps}M",
           "-bufsize", f"{mbps * 2}M", "-x264-params", "nal-hrd=cbr:force-cfr=1",
           "-c:a", "aac", "-c:s", "mov_text", "-metadata:s:s:0", "language=chi", path]
    subprocess.run(cmd, check=True)
    os.remove(srt)


SAFETY_SRT = """1
00:00:00,500 --> 00:00:04,000
欢迎参加机坪安全培训 Welcome to apron safety training

2
00:00:04,500 --> 00:00:09,000
进入机坪必须穿戴反光背心和护耳设备

3
00:00:09,500 --> 00:00:14,000
飞机发动机运转时禁止进入危险区域 Engine danger zone

4
00:00:14,500 --> 00:00:20,000
加油作业期间严禁烟火，消防车需在旁待命
"""

CAPTION_VTT = """WEBVTT

00:00:00.500 --> 00:00:05.000
这是 PowerPoint 内嵌的隐藏式字幕（WebVTT）

00:00:05.500 --> 00:00:10.000
牵引车推出前必须确认机轮挡已撤除 Remove wheel chocks before pushback
"""

LINKED_SRT = """1
00:00:01,000 --> 00:00:04,000
外链视频的同名字幕文件：行李分拣系统演示

2
00:00:04,500 --> 00:00:08,000
Baggage handling system demo
"""


def poster_png():
    img = Image.new("RGB", (640, 360), (20, 30, 60))
    d = ImageDraw.Draw(img)
    d.rectangle([260, 130, 380, 230], outline=(255, 255, 255), width=6)
    d.polygon([(300, 150), (300, 210), (350, 180)], fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def add_caption_part(pptx_path, slide_no=2):
    """在 pptx 中加入 vtt 字幕部件并与幻灯片建立关系（模拟 PowerPoint 隐藏式字幕）。"""
    tmp = pptx_path + ".tmp"
    with zipfile.ZipFile(pptx_path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"<Default Extension=\"xml\"", b"<Default Extension=\"vtt\" ContentType=\"text/vtt\"/><Default Extension=\"xml\"", 1)
            if item.filename == f"ppt/slides/_rels/slide{slide_no}.xml.rels":
                data = data.replace(b"</Relationships>",
                                    b"<Relationship Id=\"rIdCap1\" Type=\"http://schemas.microsoft.com/office/2016/11/relationships/captions\" Target=\"../media/caption1.vtt\"/></Relationships>")
            zout.writestr(item, data)
        zout.writestr("ppt/media/caption1.vtt", CAPTION_VTT.encode("utf-8"))
    os.replace(tmp, pptx_path)


def make_linked_video(pptx_path, video_name):
    """把嵌入视频的关系改为外链（TargetMode=External）。"""
    tmp = pptx_path + ".tmp"
    with zipfile.ZipFile(pptx_path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename.startswith("ppt/media/") and item.filename.endswith(".mp4"):
                continue
            data = zin.read(item.filename)
            if item.filename.startswith("ppt/slides/_rels/"):
                root = etree.fromstring(data)
                for rel in root:
                    if rel.get("Target", "").endswith(".mp4"):
                        rel.set("Target", video_name)
                        rel.set("TargetMode", "External")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    os.replace(tmp, pptx_path)


def make_convert_deck(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    prs.core_properties.author = "测试作者"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    set_bg(s, DARK)
    s.shapes.title.text = "机场运行指挥中心 AOC"
    style_text(s.shapes.title, WHITE)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(5.5), Inches(1.2))
    tb.text_frame.text = "白色正文：统一指挥、协同保障"
    p = tb.text_frame.add_paragraph()
    r = p.add_run(); r.text = "橙色强调：航班延误预警"
    r.font.color.rgb = RGBColor(0xFF, 0x99, 0x00)
    style_text(tb, WHITE)
    r.font.color.rgb = RGBColor(0xFF, 0x99, 0x00)
    for i, c in enumerate([RGBColor(0x2A, 0x4D, 0x8F), RGBColor(0xC0, 0x1E, 0x2E)]):
        box = s.shapes.add_shape(1, Inches(0.6 + i * 3), Inches(2.8), Inches(2.6), Inches(1.0))
        box.fill.solid(); box.fill.fore_color.rgb = c
        box.text_frame.text = ["运行监控", "应急处置"][i]
    logo = Image.new("RGBA", (200, 200), (0, 0, 0, 0))
    ImageDraw.Draw(logo).ellipse([10, 10, 190, 190], fill=(255, 255, 255, 255))
    buf = io.BytesIO(); logo.save(buf, "PNG"); buf.seek(0)
    s.shapes.add_picture(buf, Inches(8.8), Inches(0.2), Inches(0.9), Inches(0.9))
    add_table(s, [["指标", "数值"], ["放行正常率", "91%"]], WHITE)
    s.notes_slide.notes_text_frame.text = "用于测试黑白底互转：橙色文字应变深红，logo 应变深红，图形应变浅灰。"
    prs.save(path)


def make_word_docs(wdir):
    try:
        import docx
        from docx.shared import Inches as DInches
    except ImportError:
        print("未安装 python-docx，跳过 Word 测试文件", file=sys.stderr)
        return
    os.makedirs(wdir, exist_ok=True)
    d = docx.Document()
    d.core_properties.author = "测试作者"
    d.core_properties.title = "航班过站保障作业指导书"
    d.add_heading("航班过站保障作业指导书", 0)
    p = d.add_paragraph("本指导书规定了窄体机过站（Turnaround）保障的标准作业流程，适用于所有地面服务人员。")
    d.add_comment(p.runs, text="请补充宽体机的过站标准。", author="审核人", initials="SH")
    d.add_heading("1 适用范围 Scope", 1)
    d.add_paragraph("适用于本场所有窄体机过站航班，包括 A320 系列与 B737 系列。")
    d.add_heading("2 关键里程碑", 1)
    for t in ["AIBT 实际上轮挡时间", "TOBT 目标撤轮挡时间，由地服负责更新", "TSAT 目标启动许可时间", "AOBT 实际撤轮挡时间"]:
        d.add_paragraph(t, style="List Bullet")
    d.add_heading("2.1 TOBT 更新规则", 2)
    d.add_paragraph("TOBT 误差超过 5 分钟必须重新申报；关舱门前 15 分钟锁定。")
    d.add_heading("3 保障节点标准时间", 1)
    tb = d.add_table(rows=1, cols=3)
    tb.style = "Table Grid"
    for c, t in zip(tb.rows[0].cells, ["保障节点", "标准时长(分钟)", "责任单位"]):
        c.text = t
    for row in [("下客", "8", "机务/地服"), ("清洁", "10", "清洁队"), ("配餐", "15", "航食"), ("加油", "14", "航油"), ("上客", "15", "地服")]:
        cells = tb.add_row().cells
        for c, t in zip(cells, row):
            c.text = t
    d.add_heading("4 安全要求", 1)
    p = d.add_paragraph("加油作业期间旅客上下机须有消防车监护，严禁在危险区域内停留。")
    d.add_comment(p.runs, text="引用民航局相关规章编号。", author="安全部", initials="AQ")
    img = Image.new("RGB", (400, 120), (192, 0, 0))
    ImageDraw.Draw(img).text((20, 50), "Turnaround Flow", fill=(255, 255, 255))
    buf = io.BytesIO(); img.save(buf, "PNG"); buf.seek(0)
    d.add_picture(buf, width=DInches(4))
    d.save(os.path.join(wdir, "航班过站保障作业指导书.docx"))

    e = docx.Document()
    e.core_properties.author = "测试作者"
    e.add_heading("Smart Airport AI Platform Design", 0)
    e.add_heading("Overview", 1)
    e.add_paragraph("The platform combines computer vision on the apron with a RAG knowledge assistant for operations staff.")
    e.add_heading("Retrieval Design", 1)
    e.add_paragraph("Hybrid retrieval uses BM25 and dense embeddings with reciprocal rank fusion; a cross-encoder reranks the top results.")
    e.add_heading("Cost Optimization", 1)
    e.add_paragraph("Edge inference reduces cloud cost; FinOps dashboards track GPU usage per stand.")
    e.save(os.path.join(wdir, "Smart Airport AI Platform Design.docx"))

    if shutil.which("soffice"):
        tmpdir = os.path.join(wdir, "_tmp")
        os.makedirs(tmpdir, exist_ok=True)
        src = os.path.join(tmpdir, "旧版Word_除冰作业规范.docx")
        f = docx.Document()
        f.add_heading("除冰作业规范", 1)
        f.add_paragraph("冬季除冰作业须在除冰坪进行，除冰排队纳入 A-CDM 排序。")
        f.save(src)
        subprocess.run(["soffice", "--headless", "--convert-to", "doc", "--outdir", tmpdir, src],
                       check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        legacy = os.path.join(tmpdir, "旧版Word_除冰作业规范.doc")
        if os.path.exists(legacy):
            shutil.move(legacy, os.path.join(wdir, "旧版Word_除冰作业规范.doc"))
        shutil.rmtree(tmpdir)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out", nargs="?", default="test-ppts")
    ap.add_argument("--big-mb", type=int, default=25, help="大视频文件的目标大小(MB)")
    args = ap.parse_args()
    out = os.path.abspath(args.out)
    if os.path.exists(out):
        shutil.rmtree(out)
    os.makedirs(out)
    random.seed(7)
    j = lambda *p: os.path.join(out, *p)

    # 1. 机场运营：白底 / 黑底 / v2 副本 / A-CDM
    light = build_deck(j("机场运营", "航班过站保障Turnaround_白底.pptx"), TURNAROUND)
    build_deck(j("机场运营", "航班过站保障Turnaround_黑底.pptx"), TURNAROUND, dark=True)
    v2 = [list(s) for s in TURNAROUND]
    v2[6] = ("优化建议 Recommendations (2026 修订)", v2[6][1] + ["建立过站保障数字孪生看板"], v2[6][2])
    build_deck(j("机场运营", "航班过站保障Turnaround_v2(副本).pptx"), [tuple(s) for s in v2[:6]] + [tuple(v2[6])], hidden=(4,))
    build_deck(j("机场运营", "A-CDM机场协同决策介绍.pptx"), ACDM)

    # 2. 重复文件（名称、大小、作者、修改时间完全一致）
    os.makedirs(j("重复副本"), exist_ok=True)
    dup = j("重复副本", "航班过站保障Turnaround_白底.pptx")
    shutil.copy2(light, dup)

    # 3. 技术方案：母版级深色（英文为主） + 白底英文 + 整页图片背景
    build_deck(j("技术方案", "Smart Airport AI Platform_dark.pptx"), AI_PLATFORM, master_dark=True)
    build_deck(j("技术方案", "Smart Airport AI Platform_light.pptx"), AI_PLATFORM)
    build_deck(j("技术方案", "Cloud Architecture Overview.pptx"), CLOUD)

    def full_bg_picture(prs, fg):
        img = Image.new("RGB", (1280, 720))
        d = ImageDraw.Draw(img)
        for y in range(720):
            c = int(10 + 30 * y / 720)
            d.line([(0, y), (1280, y)], fill=(c, c, c + 15))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=80)
        buf.seek(0)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        pic = s.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
        tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1.5))
        tb.text_frame.text = "数字孪生机场 Digital Twin Airport"
        style_text(tb, WHITE, 36)
        s.notes_slide.notes_text_frame.text = "这一页用整页深色图片作为背景，用来测试图片背景的黑白底识别。"
    build_deck(j("技术方案", "数字孪生机场_图片背景.pptx"), [], extra=full_bg_picture)

    # 4. 财务（无备注页 + 无标题页）
    def untitled(prs, fg):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
        tb.text_frame.text = "附录：数据口径说明\n收入按权责发生制确认；同比口径已剔除并表范围变化影响。"
    build_deck(j("财务报告", "2026Q3经营分析.pptx"), FINANCE, extra=untitled)

    # 5. 培训视频：大文件（内封字幕视频 + vtt 隐藏式字幕）、外链视频 + srt
    has_ffmpeg = shutil.which("ffmpeg") is not None
    vdir = j("培训视频")
    os.makedirs(vdir, exist_ok=True)
    if has_ffmpeg:
        big_mp4 = os.path.join(vdir, "_big.mp4")
        make_video(big_mp4, 20, (1280, 720), max(1, args.big_mb * 8 // 20), SAFETY_SRT)

        def add_video(prs, fg):
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = "机坪安全培训视频 Apron Safety Video"
            s.shapes.add_movie(big_mp4, Inches(1), Inches(1.3), Inches(8), Inches(4), poster_frame_image=poster_png(), mime_type="video/mp4")
            s.notes_slide.notes_text_frame.text = "播放视频前提醒学员注意视频中的四条安全要点。"
        safety = [("地勤安全培训 Ground Safety Training", ["机坪安全基本要求", "危险区域识别", "应急处置流程"], "本次培训时长约 30 分钟。")]
        p = build_deck(j("培训视频", "地勤安全培训_含视频(大文件).pptx"), safety, extra=add_video)
        add_caption_part(p, slide_no=2)
        os.remove(big_mp4)

        small_mp4 = os.path.join(vdir, "行李分拣演示.mp4")
        make_video(small_mp4, 8, (320, 180), 1, LINKED_SRT)

        def add_linked(prs, fg):
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = "行李分拣系统 Baggage Handling System"
            s.shapes.add_movie(small_mp4, Inches(1), Inches(1.3), Inches(8), Inches(4), poster_frame_image=poster_png(), mime_type="video/mp4")
            s.notes_slide.notes_text_frame.text = "视频是外链文件，和 PPT 放在同一个目录下。"
        p2 = build_deck(j("培训视频", "行李系统介绍_外链视频.pptx"), [], extra=add_linked)
        make_linked_video(p2, "行李分拣演示.mp4")
        with open(os.path.join(vdir, "行李分拣演示.srt"), "w", encoding="utf-8") as f:
            f.write(LINKED_SRT)
    else:
        print("未找到 ffmpeg，跳过视频类测试文件", file=sys.stderr)

    # 6. 异常与边界文件
    edge = j("其他", "边界情况")
    os.makedirs(edge, exist_ok=True)
    Presentation().save(os.path.join(edge, "空白演示.pptx"))
    with open(os.path.join(edge, "损坏文件.pptx"), "wb") as f:
        f.write(b"PK\x03\x04" + bytes(random.getrandbits(8) for _ in range(2048)))
    with open(os.path.join(edge, "~$航班过站保障Turnaround_白底.pptx"), "wb") as f:
        f.write(b"\x05weber" + b"\x00" * 160)
    with open(os.path.join(edge, "说明.txt"), "w", encoding="utf-8") as f:
        f.write("非 PPT 文件，应被忽略\n")
    if shutil.which("soffice"):
        tmpdir = os.path.join(out, "_tmp_ppt")
        os.makedirs(tmpdir, exist_ok=True)
        src = os.path.join(tmpdir, "旧版格式_A-CDM.pptx")
        build_deck(src, ACDM[:2])
        subprocess.run(["soffice", "--headless", "--convert-to", "ppt", "--outdir", tmpdir, src],
                       check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        legacy = os.path.join(tmpdir, "旧版格式_A-CDM.ppt")
        if os.path.exists(legacy):
            shutil.move(legacy, os.path.join(j("其他"), "旧版格式_A-CDM.ppt"))
        shutil.rmtree(tmpdir)

    # 6b. 黑白底互转测试：彩色文字、小 logo、带填充的图形、表格
    make_convert_deck(j("技术方案", "黑白互转测试_黑底.pptx"))

    # 7. Word 文档（按标题分节解析、表格、列表、图片、批注、旧版 .doc）
    make_word_docs(j("Word文档"))

    # 统一修改时间，保证"重复副本"与原文件元数据一致
    st = os.stat(light)
    os.utime(dup, (st.st_atime, st.st_mtime))
    total = 0
    for root, _, files in os.walk(out):
        for fn in files:
            fp = os.path.join(root, fn)
            total += os.path.getsize(fp)
            print(f"{os.path.relpath(fp, out):60s} {os.path.getsize(fp) / 1024:10.1f} KB")
    print(f"共 {total / 1024 / 1024:.1f} MB -> {out}")


if __name__ == "__main__":
    main()
